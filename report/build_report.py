#!/usr/bin/env python3
"""Build the APA 7 scheduling report from executed local evidence."""
import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from layout import configure, paragraph, heading, page, table, note, font, caption

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / 'report' / 'Advanced_Shell_Deliverable_2_Report.docx'
TITLE = 'Advanced Shell Simulation with Integrated OS Concepts'
SUBTITLE = 'Deliverable 2 Process Scheduling'


def body(doc, text):
    return paragraph(doc, text)


def numeric_table(doc, label, title, headers, rows, widths):
    result = table(doc, label, title, headers, rows, widths)
    for row in result.rows[1:]:
        for cell in row.cells[1:]:
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return result


def listing(doc, lines):
    p = paragraph(doc, indent=False)
    p.paragraph_format.line_spacing = 1
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run('\n'.join(lines))
    font(r, size=10)
    r.font.name = 'Courier New'
    return p


def main():
    evidence = json.loads((ROOT / 'evidence' / 'scheduling_metrics.json').read_text())
    cases = evidence['cases']
    workload = evidence['workload']
    expected_workload = [('A', 6, 3, 0), ('B', 2, 1, 2),
                         ('C', 1, 1, 2), ('D', 1, 3, 3)]
    if [(p['id'], p['burst'], p['priority'], p['arrival']) for p in workload] != expected_workload:
        raise ValueError('Report analysis requires the documented four-process workload')
    expected_averages = {'rr_q1': (2.5, 5.0, 0.75), 'rr_q2': (2.5, 5.0, 1.5),
                         'rr_q4': (3.5, 6.0, 2.5), 'priority': (2.75, 5.25, 2.0)}
    for name, expected in expected_averages.items():
        actual = tuple(cases[name]['average'][key] for key in ('waiting', 'turnaround', 'response'))
        if actual != expected:
            raise ValueError(f'Evidence changed for {name}; review the narrative before building')
    doc = Document()
    configure(doc)
    doc.core_properties.title = TITLE
    doc.core_properties.subject = SUBTITLE
    p = paragraph(doc, TITLE + '\n' + SUBTITLE, indent=False, centered=True,
                  bold=True, style='Title')
    p.paragraph_format.space_before = Pt(72)
    paragraph(doc, indent=False)
    for text in ['Shobhit Dhamania', 'University of the Cumberlands',
                 'MSCS-630-M20: Advanced Operating Systems', 'Sean Perryman',
                 'October 24, 2026']:
        paragraph(doc, text, indent=False, centered=True)
    page(doc)

    heading(doc, TITLE + '\n' + SUBTITLE)
    body(doc, 'This extension adds Round-Robin and preemptive priority scheduling to the C++17 shell from Deliverable 1. The shell now controls a simulated ready queue, dispatch decisions, remaining CPU bursts, and per-process metrics. Its timed traces show that a higher-priority arrival interrupts unfinished work, while a configurable Round-Robin quantum rotates ready processes. The comparison also exposes a practical tradeoff: reducing the quantum improves initial response on this workload but increases dispatch frequency.')
    heading(doc, 'Simulation Design')
    body(doc, 'I kept the scheduler separate from the host-process job table. The inherited shell still launches external commands through fork and exec, manages process groups, and implements jobs, fg, and bg. The new sched command instead operates on simulated identifiers such as A and B. These identifiers are not operating-system PIDs, and the simulator does not change the host scheduling policy or suspend arbitrary external programs.')
    body(doc, 'Each workload entry stores an identifier, CPU burst, priority, and arrival time. Each run builds fresh remaining-time, first-dispatch, and completion records, preserving the original workload for another comparison. This separation prevents a misleading result in which the second algorithm receives processes that the first algorithm has already completed. Smaller nonnegative priority values mean higher priority throughout the interface.')
    body(doc, 'The model assumes one CPU, known CPU-only bursts, no I/O blocking, and zero modeled dispatch overhead. These are useful experimental boundaries rather than claims about a complete operating system. The distinction between a scheduling policy and its execution mechanism informed the design: the shell computes the next event, while the host provides the timer used to pace that event (Arpaci-Dusseau & Arpaci-Dusseau, 2023).')

    heading(doc, 'Queue Policies and Timed Execution')
    heading(doc, 'Round Robin', 2)
    body(doc, 'The Round-Robin ready queue is a std::deque. Dispatch removes its front entry and gives that process the smaller of its quantum and remaining burst. An unfinished process returns to the tail after quantum expiry; a completed process is removed immediately. Arrivals are admitted in chronological order, with insertion order breaking simultaneous ties. An arrival exactly at slice expiry enters before the current process is requeued. Stating this boundary rule makes the trace and metric expectations reproducible.')
    body(doc, 'The command sched run rr QUANTUM TICK_MS exposes the quantum to the user. A quantum of 2 allows process C, whose burst is 1, to finish after one unit rather than holding the CPU for an unused second unit. Later arrivals join the tail without interrupting an active Round-Robin slice. The result follows the time-sharing idea described by Arpaci-Dusseau and Arpaci-Dusseau (2023).')
    heading(doc, 'Preemptive Priority and Stable Ties', 2)
    body(doc, 'The priority policy uses std::priority_queue, a heap ordered by priority, original arrival, and insertion order. At an arrival event, the simulator compares the best ready priority with the running priority. A strictly better priority preempts unfinished work immediately at that simulated boundary. Equal priorities do not preempt. A preempted process retains its earlier FCFS rank, preventing a later peer from moving ahead merely because the older process was interrupted. POSIX likewise distinguishes priority selection from ordering within a priority list (Institute of Electrical and Electronics Engineers [IEEE] & The Open Group, 2024a).')

    heading(doc, 'Timer Semantics and Measurement')
    body(doc, 'The execution loop sleeps only until the next arrival, completion, or quantum boundary. It uses std::this_thread::sleep_for with a configurable millisecond scale; the recorded comparisons use 10 ms per simulated unit. A higher-priority arrival is therefore not delayed until the lower-priority process finishes its full burst. Idle gaps are also paced. Setting TICK_MS to 0 removes waiting for tests while retaining the same logical schedule.')
    body(doc, 'Arrivals are declared before a run and admitted during execution. The foreground simulation does not accept live shell input, so this design demonstrates timed admission from an event list rather than interactive job injection. Logical preemption is exact at the event boundary; physical wakeup can occur later because sleep resolution and host scheduling add delay. POSIX explicitly allows requested sleep intervals to overrun for those reasons (IEEE & The Open Group, 2024b).')
    numeric_table(doc, 'Table 1', 'Shared Workload for All Four Comparisons',
                  ['Process', 'Arrival', 'Burst', 'Priority'],
                  [[p['id'], str(p['arrival']), str(p['burst']), str(p['priority'])] for p in workload],
                  [1.3, 1.7, 1.7, 1.8])
    note(doc, 'Arrival and burst use simulated units. Smaller priority values represent higher priority. A, B, C, and D are simulated identifiers.')
    body(doc, 'Turnaround is completion minus arrival. Response is first dispatch minus arrival. Waiting is turnaround minus CPU burst under the stated no-I/O assumptions; it includes ready time after preemption. For example, priority process A completes at 9 after arriving at 0 with burst 6, so its turnaround is 9 and waiting is 3. Its response remains 0 because it first ran immediately. Waiting and response therefore answer different questions.')

    heading(doc, 'Performance Results')
    labels = [('rr_q1', 'Round-Robin q = 1'), ('rr_q2', 'Round-Robin q = 2'),
              ('rr_q4', 'Round-Robin q = 4'), ('priority', 'Priority')]
    numeric_table(doc, 'Table 2', 'Average Scheduling Metrics and Dispatch Counts',
                  ['Policy', 'Waiting', 'Turnaround', 'Response', 'Dispatches'],
                  [[label, f"{cases[key]['average']['waiting']:.2f}",
                    f"{cases[key]['average']['turnaround']:.2f}",
                    f"{cases[key]['average']['response']:.2f}", str(cases[key]['dispatches'])]
                   for key, label in labels], [2.1, 1.0, 1.3, 1.1, 1.0])
    note(doc, 'Time metrics use simulated units and are arithmetic means of four processes. Every policy completes the total burst of 10 units. Dispatch counts include assigning the CPU again to the same process.')
    body(doc, 'Changing the Round-Robin quantum from 1 to 2 doubles average response from 0.75 to 1.50 while leaving average waiting at 2.50 and turnaround at 5.00. Dispatches fall from 10 to 6. A quantum of 4 reduces dispatches to 5, but average response rises to 2.50 and waiting to 3.50. Thus fewer dispatches are not sufficient evidence of a better experience for processes waiting to start.')
    body(doc, 'These results also limit a common shortcut in interpretation. Quantum 1 does not improve every metric relative to quantum 2: their average waiting and turnaround are identical for this workload. The smaller quantum changes who receives early CPU service and increases scheduling events. On a real CPU, each switch has costs that this logical model omits; the dispatch count reveals exposure to such costs without pretending to measure them. The advantage of a chosen quantum remains workload dependent.')
    body(doc, 'Priority produces average waiting 2.75, turnaround 5.25, and response 2.00. It favors urgent work but does not minimize these aggregate metrics in this comparison. The process-level results explain why: B starts at arrival, while equal low-priority process D waits substantially longer. Choosing a policy therefore requires an objective, such as urgent service, responsiveness, or fairness, rather than a single claim that one algorithm is always faster.')

    heading(doc, 'Preemption and Process Level Effects')
    numeric_table(doc, 'Table 3', 'Per Process Metrics Under Preemptive Priority',
                  ['Process', 'First run', 'Completion', 'Waiting', 'Turnaround', 'Response'],
                  [[m['id'], str(m['first_start']), str(m['completion']), str(m['waiting']),
                    str(m['turnaround']), str(m['response'])] for m in cases['priority']['metrics']],
                  [0.9, 1.0, 1.3, 1.0, 1.3, 1.0])
    note(doc, 'All values use simulated units. The single observed preemption occurs at t = 2 when B interrupts A. C shares B\'s priority and retains FCFS order behind B.')
    body(doc, 'The recorded priority timeline is A from 0 to 2, B from 2 to 4, C from 4 to 5, A from 5 to 9, and D from 9 to 10. B and C arrive at 2, yet only B is selected first because it was inserted before C. D arrives at 3 with the same priority as A. When urgent work finishes at 5, A resumes before D because A originally arrived at 0. This makes the FCFS tie rule observable even after a preemption.')
    body(doc, 'D illustrates the cost of static priorities. Its burst is only 1, but its waiting and response are both 6 under priority scheduling; Round-Robin quantum 2 gives it waiting and response of 4. The finite event list still completes every process. An unbounded stream of better-priority arrivals could starve low-priority work, so a production extension could add aging or service guarantees. Those mechanisms are improvements, not behaviors claimed by this implementation.')
    elapsed = ', '.join(f"{label}: {cases[key]['wall_ms']:.3f} ms" for key, label in labels)
    body(doc, 'The recorded elapsed wall times were ' + elapsed + '. These are single-run observations from the evidence files, not repeated-trial performance estimates. Each run requests approximately 100 ms of sleep for 10 CPU units. Printing, timer overshoot, and host scheduling explain additional elapsed time. The few-millisecond differences do not establish that one policy executes faster on hardware.')

    heading(doc, 'Verification and Design Lessons')
    body(doc, 'The build enables C++17 with warnings treated as errors. Verification retains the parser assertions and 20 inherited shell integration tests, then adds scheduler unit checks and four scheduling integration tests. An independent unit-at-a-time reference model evaluates 40 generated workloads under both algorithms, giving 80 schedule comparisons. It checks completion, first response, waiting, and turnaround rather than accepting a plausible event label as sufficient evidence.')
    body(doc, 'The focused cases exercise higher-priority preemption, equal-priority stability, arrivals exactly at quantum boundaries, short final slices, initially idle CPUs, repeated runs, and invalid input. A timed unit test checks that real waiting occurs. The scheduler unit suite also passes AddressSanitizer and UndefinedBehaviorSanitizer, and static analysis reports no scheduler findings. These checks support the tested behavior; they do not establish correctness for every possible workload or an untested Linux installation.')
    body(doc, 'I treated arrival ordering and metric accounting as the main design challenges. A heap that uses priority alone leaves equal-priority results unstable. Reassigning a new FCFS rank after preemption would wrongly let D overtake A. Keeping original arrival and insertion rank fixes that policy error. Similarly, updating the first-dispatch value on every resume would overstate response time, so the implementation records it only once.')
    body(doc, 'The useful insight from the results is that an execution trace and a performance metric reveal different parts of the same decision. At t = 2, the trace proves that B preempts A; D\'s waiting time later exposes the consequence of favoring urgent work. Separating these observations made the policy easier to evaluate without reducing the analysis to successful execution alone.')
    body(doc, 'Next steps would include live arrival injection through a synchronized event channel, explicit context-switch costs, I/O blocking, and aging. Resource limits and numeric validation already bound workloads and reject invalid changes without replacing valid state. Scheduling commands reject background syntax because parent-owned workload changes would otherwise occur in a child copy. The completed extension demonstrates both required policies while preserving the original shell interface.')
    page(doc)

    heading(doc, 'References')
    refs = [
        [('Arpaci-Dusseau, R. H., & Arpaci-Dusseau, A. C. (2023). ', False),
         ('Operating systems: Three easy pieces', True),
         (' (Version 1.10). Arpaci-Dusseau Books. https://pages.cs.wisc.edu/~remzi/OSTEP/', False)],
        [('Institute of Electrical and Electronics Engineers & The Open Group. (2024a). General information. In ', False),
         ('The Open Group base specifications issue 8', True),
         ('. https://pubs.opengroup.org/onlinepubs/9799919799/functions/V2_chap02.html', False)],
        [('Institute of Electrical and Electronics Engineers & The Open Group. (2024b). nanosleep. In ', False),
         ('The Open Group base specifications issue 8', True),
         ('. https://pubs.opengroup.org/onlinepubs/9799919799/functions/nanosleep.html', False)],
    ]
    for parts in refs:
        p = paragraph(doc, indent=False)
        p.paragraph_format.left_indent = Inches(.5)
        p.paragraph_format.first_line_indent = Inches(-.5)
        for text, italic in parts:
            font(p.add_run(text), italic=italic)
    page(doc)

    heading(doc, 'Appendix A')
    heading(doc, 'Reproducing the Recorded Scheduling Evidence')
    body(doc, 'Build with make and start ./build/os-shell. Enter the workload below once, then run each comparison. The same workload persists between runs; sched reset clears it. Full executed output and structured measurements are included under evidence/. The listing below is a selected text excerpt from the priority run, not a Terminal screenshot.')
    caption(doc, 'Listing A1', 'Workload and Comparison Commands')
    listing(doc, ['sched add A 6 3 0', 'sched add B 2 1 2', 'sched add C 1 1 2',
                  'sched add D 1 3 3', 'sched list', 'sched run rr 1 10',
                  'sched run rr 2 10', 'sched run rr 4 10', 'sched run priority 10'])
    caption(doc, 'Listing A2', 'Selected Events From the Executed Priority Trace')
    source_lines = (ROOT / 'evidence' / 'priority.txt').read_text().splitlines()
    chosen = ['t=0 DISPATCH', 't=2 ARRIVE', 't=2 PREEMPT', 't=2 DISPATCH',
              't=4 DISPATCH', 't=5 DISPATCH', 't=9 DISPATCH', 'AVERAGE']
    listing(doc, [line for line in source_lines if any(line.startswith(key) for key in chosen)])
    body(doc, 'The files examples/rr_q1.commands, rr_q2.commands, rr_q4.commands, and priority.commands also run in batch mode. scripts/collect_scheduling_evidence.py rebuilds the text records and scheduling_metrics.json from actual shell execution. Screenshot capture steps are provided in screenshots/CAPTURE_INSTRUCTIONS.txt.')

    screenshot_titles = {'01_round_robin.png': 'Round Robin With Configurable Quantum',
                         '02_priority_preemption.png': 'Priority Arrival and Preemption',
                         '03_rr_quantum1_metrics.png': 'Round Robin Quantum One Metrics',
                         '04_rr_quantum4_metrics.png': 'Round Robin Quantum Four Metrics'}
    existing = [(p.name, screenshot_titles.get(p.name, p.stem.replace('_', ' ').title()))
                for p in sorted((ROOT / 'screenshots').glob('*.png')) if p.is_file()]
    if existing:
        page(doc)
        heading(doc, 'Appendix B')
        heading(doc, 'Terminal Demonstrations')
        for i, (name, title) in enumerate(existing, 1):
            if i > 1: page(doc)
            caption(doc, f'Figure B{i}', title)
            p = paragraph(doc, indent=False)
            p.paragraph_format.line_spacing = 1
            pic = p.add_run().add_picture(str(ROOT / 'screenshots' / name), width=Inches(6.5))
            pic._inline.docPr.set('descr', title + ' captured in an actual Terminal session.')
            screenshot_notes = {
                '01_round_robin.png': 'Actual Terminal capture with quantum 2 and 10 ms per simulated unit. Tables summarize separate automated runs, so wall times may differ.',
                '02_priority_preemption.png': 'The final run uses 10 ms per simulated unit and shows priority preemption and FCFS ties. The earlier 1 ms/unit run and input correction are retained.',
                '03_rr_quantum1_metrics.png': 'Quantum 1 in fast mode (tick_ms=0). Logical metrics agree with Table 2; elapsed time is not comparable with the timed 10 ms runs.',
                '04_rr_quantum4_metrics.png': 'Quantum 4 in fast mode (tick_ms=0). Logical metrics agree with Table 2; elapsed time is not comparable with the timed 10 ms runs.'
            }
            note(doc, screenshot_notes.get(name, 'Actual Terminal capture of the shell scheduling demonstration.'))
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == '__main__':
    main()
